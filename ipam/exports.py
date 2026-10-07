"""Jadvallarni CSV, Word (.docx) va PDF formatida yuklab berish."""

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from django.http import HttpResponse
from django.utils import timezone

EXPORT_FORMATS = ("csv", "docx", "pdf")


@dataclass
class Report:
    title: str
    headers: list
    rows: list
    filename: str = "hisobot"
    info: list = field(default_factory=list)  # sarlavha ostidagi qatorlar: [("Viloyat", "Samarqand"), ...]
    totals: list | None = None  # jadval oxiridagi "Jami" qatori
    landscape: bool = True


def export_response(report, fmt):
    stamp = timezone.localtime().strftime("%Y-%m-%d_%H-%M")
    filename = f"{report.filename}_{stamp}.{fmt}"
    builders = {
        "csv": ("text/csv; charset=utf-8", build_csv),
        "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", build_docx),
        "pdf": ("application/pdf", build_pdf),
    }
    content_type, builder = builders[fmt]
    response = HttpResponse(builder(report), content_type=content_type)

    from .audit import log_event
    from .models import AuditLog

    details = {"Format": ["", fmt.upper()], "Qatorlar soni": ["", str(len(report.rows))]}
    for label, value in report.info:
        details[str(label)] = ["", str(value)]
    log_event(AuditLog.ACTION_EXPORT, object_repr=report.title, model_name="Hisobot", changes=details)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _cell(value):
    return "" if value is None else str(value)


def _generated_line():
    return f"Tuzilgan sana: {timezone.localtime():%d.%m.%Y %H:%M}"


# ---------- CSV ----------

def build_csv(report):
    buffer = io.StringIO()
    buffer.write("\ufeff")  # Excel o'zbekcha harflarni to'g'ri ko'rsatishi uchun
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(report.headers)
    for row in report.rows:
        writer.writerow([_cell(v) for v in row])
    if report.totals:
        writer.writerow([_cell(v) for v in report.totals])
    return buffer.getvalue().encode("utf-8")


# ---------- Word ----------

def build_docx(report):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    section = doc.sections[0]
    if report.landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))

    doc.add_heading(report.title, level=1)
    for label, value in report.info:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(f"{label}: ").bold = True
        p.add_run(_cell(value))
    meta = doc.add_paragraph(_generated_line())
    meta.runs[0].font.size = Pt(8)
    meta.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    table = doc.add_table(rows=1, cols=len(report.headers))
    table.style = "Table Grid"

    def shade(cell, color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    def fill(cells, values, bold=False, size=9):
        for cell, value in zip(cells, values):
            cell.text = ""
            run = cell.paragraphs[0].add_run(_cell(value))
            run.font.size = Pt(size)
            run.bold = bold

    fill(table.rows[0].cells, report.headers, bold=True)
    for cell in table.rows[0].cells:
        shade(cell, "E5E7EB")
    # Sarlavha har sahifada takrorlansin
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    tr_pr.append(header_flag)

    for row in report.rows:
        fill(table.add_row().cells, row)
    if report.totals:
        cells = table.add_row().cells
        fill(cells, report.totals, bold=True)
        for cell in cells:
            shade(cell, "F3F4F6")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------

_FONT = None


def _register_fonts():
    """O'zbekcha harflar (o', g', ʻ) chiqishi uchun Unicode shrift ulanadi."""
    global _FONT
    if _FONT:
        return _FONT
    from reportlab import rl_config
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    ]
    vera_dir = Path(rl_config.__file__).parent / "fonts"
    candidates.append((vera_dir / "Vera.ttf", vera_dir / "VeraBd.ttf"))
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("IPAM", str(regular)))
            pdfmetrics.registerFont(TTFont("IPAM-Bold", str(bold)))
            pdfmetrics.registerFontFamily("IPAM", normal="IPAM", bold="IPAM-Bold")
            _FONT = "IPAM"
            return _FONT
    _FONT = "Helvetica"
    return _FONT


def _column_widths(report, available, font, bold, size=8, padding=15, max_pref=220):
    # padding: katakning chap+o'ng ichki chekkasi (reportlab'da 6+6) va ozgina zaxira
    """Har bir ustunga: kamida eng uzun so'z sig'adigan, imkon bo'lsa butun matn
    bir qatorga sig'adigan kenglik beriladi. Joy yetmasa, uzun matnli ustunlar qisqaradi."""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    min_w, pref_w = [], []
    for i, header in enumerate(report.headers):
        cells = [(_cell(header), bold)] + [(_cell(row[i]), font) for row in report.rows[:300]]
        if report.totals:
            cells.append((_cell(report.totals[i]), bold))
        # Sarlavhadagi so'zlar bo'linishi mumkin, qiymatlardagi so'zlar (sana, summa) — yo'q
        words = [w for text, f in cells[1:] for w in text.split(" ")]
        longest_word = max((stringWidth(w, font, size) for w in words), default=0)
        header_word = max((stringWidth(w, bold, size) for w in _cell(header).split()), default=0)
        full = max(stringWidth(text, f, size) for text, f in cells)
        lo = max(longest_word, header_word) + padding
        min_w.append(lo)
        pref_w.append(max(lo, min(full + padding, max_pref)))

    if sum(pref_w) <= available:
        extra = available - sum(pref_w)
        return [w + extra * w / sum(pref_w) for w in pref_w]
    if sum(min_w) >= available:
        return [available * w / sum(min_w) for w in min_w]
    flexible = [p - m for p, m in zip(pref_w, min_w)]
    room = available - sum(min_w)
    return [m + room * f / sum(flexible) for m, f in zip(min_w, flexible)]


def build_pdf(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    font = _register_fonts()
    bold = f"{font}-Bold" if font == "IPAM" else "Helvetica-Bold"
    pagesize = landscape(A4) if report.landscape else A4

    title_style = ParagraphStyle("title", fontName=bold, fontSize=15, leading=19, spaceAfter=6)
    info_style = ParagraphStyle("info", fontName=font, fontSize=9.5, leading=13)
    meta_style = ParagraphStyle("meta", fontName=font, fontSize=8, leading=10, textColor=colors.HexColor("#6B7280"))
    cell_style = ParagraphStyle("cell", fontName=font, fontSize=8, leading=10)
    head_style = ParagraphStyle("head", parent=cell_style, fontName=bold)

    def p(value, style=cell_style):
        return Paragraph(escape(_cell(value)), style)

    story = [Paragraph(escape(report.title), title_style)]
    for label, value in report.info:
        story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(_cell(value))}", info_style))
    story += [Paragraph(escape(_generated_line()), meta_style), Spacer(1, 0.4 * cm)]

    data = [[p(h, head_style) for h in report.headers]]
    data += [[p(v) for v in row] for row in report.rows]
    if report.totals:
        data.append([p(v, head_style) for v in report.totals])

    col_widths = _column_widths(report, pagesize[0] - 3 * cm, font, bold)

    table = Table(data, colWidths=col_widths, repeatRows=1)
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CA3AF")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if report.totals:
        commands.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F3F4F6")))
    table.setStyle(TableStyle(commands))
    story.append(table)

    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.drawRightString(pagesize[0] - 1.5 * cm, 1 * cm, f"{doc.page}-sahifa")
        canvas.restoreState()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=pagesize, title=report.title,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return buffer.getvalue()


class ExportMixin:
    """Ro'yxat sahifalariga ?export=csv|docx|pdf qo'shadi.

    Ko'rinish `build_report(queryset)` metodini yozishi kerak — u filtrlangan
    (sahifalarga bo'linmagan) ma'lumotdan Report qaytaradi.
    """

    def get(self, request, *args, **kwargs):
        fmt = request.GET.get("export")
        if fmt in EXPORT_FORMATS:
            return export_response(self.build_report(self.get_queryset()), fmt)
        return super().get(request, *args, **kwargs)

    def build_report(self, queryset):
        raise NotImplementedError

    def filter_info(self):
        """Tanlangan filtrlar hisobot sarlavhasi ostiga yoziladi."""
        form = getattr(self, "filter_form", None)
        if not form or not form.is_bound or not form.is_valid():
            return []
        info = []
        for name, value in form.cleaned_data.items():
            if value in (None, ""):
                continue
            field = form.fields[name]
            if hasattr(field, "choices") and not hasattr(value, "pk"):
                value = dict(field.choices).get(value, value)
            info.append((field.label, value))
        return info
